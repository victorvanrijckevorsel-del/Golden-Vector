"""Bulk-import manual mining inputs + source verification from the friend's
Excel workbook into our local SQLite manual-data store.

Covers 56 active Tool B tickers that were blank plus re-confirms the 3 we
already had (AEM/KGC/NEM) without downgrading existing VERIFIED flags.

## What gets written

For each ticker in Excel `Screening Data`:
- `company_inputs` row with the 11 manual mining fields (production_oz,
  aisc_usd_per_oz, cash_cost_usd_per_oz, royalty_rate, sustaining_capex_musd,
  da_musd, interest_expense_musd, tax_rate, reserve_life_years,
  net_debt_musd, ebitda_ltm_musd) — raw values; `manual_store` already
  normalizes percent-stored rates (>1 → /100) and decimal-stored rates are
  left alone, so the six decimal-convention Excel cells (AGI, BGL.AX, DRD,
  EMR.AX, HOC.L, OBM.AX) land at the correct effective rate either way.

- `source_verification` rows — one per imported field, but ONLY when no
  existing verification row exists for that (ticker, field). This prevents
  the backfill from downgrading AEM/KGC/NEM rows that are already
  per-field VERIFIED. Newly-imported tickers get:
  - `aisc_usd_per_oz` → status from Excel `AISC Verification` (uppercased)
  - `production_oz` → status from Excel `Production verification`
  - all 9 other fields → `ESTIMATED` with a provenance note pointing to
    the workbook (the friend didn't separately verify those; being honest
    beats overclaiming per the cross-check review)

## What gets skipped

- Tickers inactive in our universe (e.g. NGD — Yahoo history is broken)
- Legacy fixture tickers GOLD/FNV/FRES.L (not in friend's universe)
- Friend ticker → our ticker mapping from `scripts.friend_excel_import`:
  ARMN → ARIS.TO, RMS → RMS.AX

## Usage

    python -m scripts.backfill_manual_data_from_excel           # dry-run
    python -m scripts.backfill_manual_data_from_excel --apply   # write
"""

from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import pandas as pd

from golden_vector.app.config import load_app_config
from golden_vector.app.paths import ProjectPaths
from golden_vector.screening.manual_store import (
    load_store_tables,
    upsert_company_input,
    upsert_source_verification,
)
from scripts.friend_excel_import import default_workbook_path, map_friend_ticker


# Column positions in Excel `Screening Data` (1-indexed to match openpyxl).
_SD_COL = {
    "ticker": 1,
    "production_oz": 9,
    "aisc_usd_per_oz": 10,
    "cash_cost_usd_per_oz": 11,
    "royalty_rate": 12,
    "sustaining_capex_musd": 13,
    "da_musd": 14,
    "interest_expense_musd": 15,
    "tax_rate": 16,
    "reserve_life_years": 17,
    "net_debt_musd": 18,
    "ebitda_ltm_musd": 19,
}

_MANUAL_FIELDS = [
    "production_oz",
    "aisc_usd_per_oz",
    "cash_cost_usd_per_oz",
    "royalty_rate",
    "sustaining_capex_musd",
    "da_musd",
    "interest_expense_musd",
    "tax_rate",
    "reserve_life_years",
    "net_debt_musd",
    "ebitda_ltm_musd",
]

_PROVENANCE_NOTE = "Imported from Gold_Mining_Screening_v10226_EVEB.xlsx (bulk backfill 2026-04-24)"


@dataclass(frozen=True)
class PerTickerImport:
    friend_ticker: str
    our_ticker: str
    values: dict[str, float]
    skipped_fields: list[tuple[str, str]]  # (field, reason) for cells we couldn't import
    aisc_status: str | None  # "VERIFIED" / "ESTIMATED" / None if unknown
    aisc_source_date: str | None
    aisc_source_url: str | None
    production_status: str | None
    production_source_url: str | None


def _coerce_numeric(value: object) -> tuple[float | None, str | None]:
    """Try to coerce a cell value to float. Returns (value, skip_reason)."""
    if value is None:
        return None, "empty"
    if isinstance(value, str):
        stripped = value.strip()
        if not stripped or stripped.startswith("#") or stripped.upper() == "N/A":
            return None, f"cell error: {stripped!r}"
        try:
            return float(stripped), None
        except ValueError:
            return None, f"non-numeric: {stripped!r}"
    try:
        return float(value), None
    except (TypeError, ValueError):
        return None, f"non-numeric: {value!r}"


def _coerce_text(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def _coerce_iso_date(value: object) -> tuple[str | None, str | None]:
    """Try to coerce a cell value to YYYY-MM-DD. Free-form strings like
    "Q3 2025" or "FY26 Guide" can't be parsed — return (None, original)
    so the caller can stash the original text in a notes field.
    """
    if value is None:
        return None, None
    if hasattr(value, "isoformat"):
        # openpyxl returns datetime.date / datetime.datetime for date cells
        try:
            return value.isoformat()[:10], None
        except Exception:
            return None, str(value)
    text = str(value).strip()
    if not text:
        return None, None
    try:
        return str(pd.to_datetime(text, errors="raise").date()), None
    except Exception:
        # Unparseable like "2025 Guide", "Q3 2025", "FY26". Keep original
        # text so the caller can put it in notes.
        return None, text


def read_screening_data(workbook_path: Path) -> list[PerTickerImport]:
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        ws = wb["Screening Data"]
        imports: list[PerTickerImport] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            friend_ticker = row[_SD_COL["ticker"] - 1] if row else None
            if not friend_ticker:
                continue
            friend_ticker = str(friend_ticker).strip()
            our_ticker = map_friend_ticker(friend_ticker)

            values: dict[str, float] = {}
            skipped: list[tuple[str, str]] = []
            for field in _MANUAL_FIELDS:
                col_idx = _SD_COL[field] - 1
                raw = row[col_idx] if col_idx < len(row) else None
                numeric, err = _coerce_numeric(raw)
                if numeric is None:
                    if err and err != "empty":
                        skipped.append((field, err))
                    continue
                values[field] = numeric

            imports.append(PerTickerImport(
                friend_ticker=friend_ticker,
                our_ticker=our_ticker,
                values=values,
                skipped_fields=skipped,
                aisc_status=None,
                aisc_source_date=None,
                aisc_source_url=None,
                production_status=None,
                production_source_url=None,
            ))
        return imports
    finally:
        wb.close()


def read_verification_sheets(
    workbook_path: Path,
) -> tuple[dict[str, dict], dict[str, dict]]:
    """Return (aisc_verification, production_verification) keyed by friend ticker."""
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        aisc: dict[str, dict] = {}
        ws = wb["AISC Verification"]
        # Row 1 is header: Ticker | Company | Source Date | Verified AISC | Status | Notes | Verified Source | Source URL
        for row in ws.iter_rows(min_row=2, values_only=True):
            ticker = row[0] if row else None
            if not ticker:
                continue
            aisc[str(ticker).strip()] = {
                "status": _coerce_text(row[4]),
                "source_date": _coerce_text(row[2]),
                "source_url": _coerce_text(row[7]) if len(row) > 7 else None,
                "notes": _coerce_text(row[5]) if len(row) > 5 else None,
            }

        production: dict[str, dict] = {}
        ws = wb["Production verification"]
        # Row 4 is header: Ticker | Company | Original | Final 2026 | Change % | Status | Source/Notes | Source URL
        for row in ws.iter_rows(min_row=5, values_only=True):
            ticker = row[0] if row else None
            if not ticker:
                continue
            production[str(ticker).strip()] = {
                "status": _coerce_text(row[5]),
                "notes": _coerce_text(row[6]) if len(row) > 6 else None,
                "source_url": _coerce_text(row[7]) if len(row) > 7 else None,
            }
        return aisc, production
    finally:
        wb.close()


def _enrich_with_verification(
    imports: list[PerTickerImport],
    aisc: dict[str, dict],
    production: dict[str, dict],
) -> list[PerTickerImport]:
    enriched: list[PerTickerImport] = []
    for item in imports:
        aisc_row = aisc.get(item.friend_ticker, {})
        prod_row = production.get(item.friend_ticker, {})
        enriched.append(PerTickerImport(
            friend_ticker=item.friend_ticker,
            our_ticker=item.our_ticker,
            values=item.values,
            skipped_fields=item.skipped_fields,
            aisc_status=(aisc_row.get("status") or "").upper() or None,
            aisc_source_date=aisc_row.get("source_date"),
            aisc_source_url=aisc_row.get("source_url"),
            production_status=(prod_row.get("status") or "").upper() or None,
            production_source_url=prod_row.get("source_url"),
        ))
    return enriched


def _active_ticker_set(paths: ProjectPaths) -> set[str]:
    loaded = load_app_config(paths)
    return {
        t.ticker for t in loaded.app.universe.tickers
        if t.active and t.tool_b_enabled
    }


def _existing_verification_index(paths: ProjectPaths) -> set[tuple[str, str]]:
    """Return set of (ticker, field_name) that already have a verification row."""
    _, source_verification, _, _ = load_store_tables(paths)
    if source_verification.empty:
        return set()
    return {
        (str(row["ticker"]).upper(), str(row["field_name"]))
        for _, row in source_verification.iterrows()
    }


def print_preview(imports: list[PerTickerImport], active_tickers: set[str]) -> None:
    in_scope = [i for i in imports if i.our_ticker in active_tickers]
    out_of_scope = [i for i in imports if i.our_ticker not in active_tickers]

    print(f"Friend tickers in workbook: {len(imports)}")
    print(f"  -> will import: {len(in_scope)} (active + tool_b_enabled)")
    print(f"  -> will skip:   {len(out_of_scope)} (inactive or absent from universe)")
    if out_of_scope:
        for item in out_of_scope:
            reason = "not in universe" if item.our_ticker == item.friend_ticker else f"mapped to {item.our_ticker}, not active"
            if item.our_ticker in {"NGD"}:
                reason = "NGD inactive (Yahoo history broken)"
            print(f"     - {item.friend_ticker} ({reason})")

    print()
    print("First 3 tickers' extracted data:")
    for item in in_scope[:3]:
        print(f"  {item.friend_ticker:10} -> {item.our_ticker}")
        for field, value in item.values.items():
            print(f"    {field:28} = {value}")
        if item.skipped_fields:
            for field, reason in item.skipped_fields:
                print(f"    SKIP {field:28}   ({reason})")
        print(f"    AISC status: {item.aisc_status}, Production status: {item.production_status}")


def apply_backfill(
    paths: ProjectPaths,
    imports: list[PerTickerImport],
    active_tickers: set[str],
    existing_verification: set[tuple[str, str]],
) -> dict[str, int]:
    stats = {
        "tickers_imported": 0,
        "tickers_skipped_inactive": 0,
        "fields_written": 0,
        "fields_skipped_cell_error": 0,
        "verification_rows_written": 0,
        "verification_rows_skipped_existing": 0,
    }

    for item in imports:
        if item.our_ticker not in active_tickers:
            stats["tickers_skipped_inactive"] += 1
            continue

        # company_inputs upsert (one call with all available fields)
        if item.values:
            upsert_company_input(paths, ticker=item.our_ticker, values=item.values)
            stats["fields_written"] += len(item.values)
        stats["fields_skipped_cell_error"] += len(item.skipped_fields)
        stats["tickers_imported"] += 1

        # source_verification upserts — one per field, only if not already present
        for field in item.values:
            key = (item.our_ticker, field)
            if key in existing_verification:
                stats["verification_rows_skipped_existing"] += 1
                continue

            if field == "aisc_usd_per_oz":
                status = item.aisc_status or "ESTIMATED"
                iso_date, raw_date_text = _coerce_iso_date(item.aisc_source_date)
                # If the friend's source_date is free-form (e.g. "Q3 2025"),
                # fold the original text into notes so we don't lose it.
                notes_parts = [_PROVENANCE_NOTE]
                if raw_date_text:
                    notes_parts.append(f"Source date from workbook: {raw_date_text}")
                verification_values = {
                    "source_date": iso_date,
                    "source_url": item.aisc_source_url,
                    "notes": "; ".join(notes_parts),
                }
            elif field == "production_oz":
                status = item.production_status or "ESTIMATED"
                verification_values = {
                    "source_url": item.production_source_url,
                    "notes": _PROVENANCE_NOTE,
                }
            else:
                status = "ESTIMATED"
                verification_values = {"notes": _PROVENANCE_NOTE}

            # Drop None values (upsert expects real values or skip)
            verification_values = {k: v for k, v in verification_values.items() if v is not None}

            upsert_source_verification(
                paths,
                ticker=item.our_ticker,
                field_name=field,
                verification_status=status,
                values=verification_values,
            )
            stats["verification_rows_written"] += 1

    return stats


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--apply",
        action="store_true",
        help="Write to SQLite. Default is dry-run (prints preview only).",
    )
    args = parser.parse_args()

    paths = ProjectPaths.discover()
    workbook = default_workbook_path(paths.repo_root)
    if not workbook.exists():
        print(f"ERROR: workbook not found at {workbook}", file=sys.stderr)
        return 1

    imports = read_screening_data(workbook)
    aisc_sheet, production_sheet = read_verification_sheets(workbook)
    imports = _enrich_with_verification(imports, aisc_sheet, production_sheet)

    active_tickers = _active_ticker_set(paths)
    existing_verification = _existing_verification_index(paths)

    print_preview(imports, active_tickers)
    print()
    print(f"Existing (ticker,field) verification rows: {len(existing_verification)}")
    print("  (these will NOT be overwritten — protects AEM/KGC/NEM VERIFIED state)")

    if not args.apply:
        print()
        print("(dry-run — re-run with --apply to write)")
        return 0

    stats = apply_backfill(paths, imports, active_tickers, existing_verification)
    print()
    print("=== Backfill applied ===")
    for key, count in stats.items():
        print(f"  {key}: {count}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
